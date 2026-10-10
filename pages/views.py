"""
views.py — صفحة "عن الإدارة" (Home) + "قبول المنح" + "اتصل بنا"
هذا الملف تجربة تحويل، لإثبات جدوى الخيار الثالث (قوالب Django + Alpine).
البيانات هنا (partners, goals, news, categories, benefits...) ثوابت في الكود كما هو مقرر في قرارات المشروع.
"""
from django.shortcuts import render
from django.templatetags.static import static
from django.core.mail import EmailMultiAlternatives
from django.template.loader import render_to_string
from django.http import JsonResponse
from django.conf import settings
from django.utils import timezone
import os
from openpyxl import Workbook, load_workbook


PROGRAM_STEPS = [
    {"title": "تبدأ", "desc": "تقديم طلب الرغبة في الدعم عبر قنوات التواصل الخاصة بالإدارة",
     "color": "#B3DEC7", "cx": 563, "cy": 117, "lx": 483, "ly": 4, "lw": 160, "lh": 60},
    {"title": "تختار", "desc": "اختيار واعتماد المشروع بعد عرض عدد من المشاريع المقترحة",
     "color": "#5AB379", "cx": 650, "cy": 290, "lx": 700, "ly": 252, "lw": 160, "lh": 76},
    {"title": "تمنح", "desc": "إصدار قرار قبول المنحة والتوقيع رسمياً والإعلان عنها",
     "color": "#31715F", "cx": 563, "cy": 463, "lx": 483, "ly": 516, "lw": 160, "lh": 62},
    {"title": "تؤثر", "desc": "عكس أثر المنحة من خلال تقارير دورية أثر وقصص النجاح",
     "color": "#6B7787", "cx": 337, "cy": 463, "lx": 257, "ly": 516, "lw": 160, "lh": 62},
    {"title": "تستفيد", "desc": "الاستفادة من المزايا المقدمة لتعظيم الأثر وتحقيق القيمة",
     "color": "#17343D", "cx": 250, "cy": 290, "lx": 40, "ly": 252, "lw": 160, "lh": 76},
    {"title": "تكرم", "desc": "تكرم مساهمة المانح تقديراً لعطائه وأثره",
     "color": "#D7CF60", "cx": 337, "cy": 117, "lx": 257, "ly": 4, "lw": 160, "lh": 60},
]

GOALS = [
    "بناء وتعزيز العلاقات مع المانحين بما يحقق طموحاتهم ويتواءم مع استراتيجيات الجامعة.",
    "المساهمة في تحقيق أهداف رؤية المملكة 2030 من خلال الشراكة مع القطاعين العام والخاص.",
    "المساهمة في خدمة احتياجات المجتمع بإبراز الأثر الناتج عن العلاقة المتبادلة بين الجامعة ومانحيها.",
    "تحقيق الاستدامة المالية وتعزيز انخراط الشركاء والمانحين في خدمة المجتمع.",
]

NEWS_ITEMS = [
    {
        "img": static("images/news-bandah.jpg"),
        "tag": "شراكة", "date": "17 يونيو 2025م",
        "title": "منحة بندة للتجزئة لدعم عيادات الأسنان المتنقلة",
        "desc": "وقّعت الجامعة اتفاقية شراكة مع بندة للتجزئة منحة لدعم عيادات الأسنان المتنقلة، لما لها من أثر في تحسين جودة الرعاية الصحية المقدمة للمجتمع، وتعزيز التدريب العملي لطلبة كلية طب الأسنان.",
        "link": "https://x.com/kauedu_sa/status/1935008353710649599?s=48&t=G7Y-7ksNJMZzV3q9M9rCqg",
    },
    {
        "img": static("images/news-aqeel.png"),
        "tag": "شراكة", "date": "19 يونيو 2025م",
        "title": "منحة بحثية من مؤسسة العقيل الخيرية لدعم البحث العلمي",
        "desc": 'وقّعت الجامعة اتفاقية تعاون مع مؤسسة أبناء عبدالرحمن العقيل الخيرية "جرير للاستثمار" لتقديم منحة بحثية تمتد لخمسة أعوام، لدعم منظومة البحث العلمي وتمكين الباحثين من المشاركة العالمية.',
        "link": "https://x.com/kauedu_sa/status/1935699904321519842?s=48&t=G7Y-7ksNJMZzV3q9M9rCqg",
    },
    {
        "img": static("images/news-kidney-dialysis.jpg"),
        "tag": "شراكة", "date": "",
        "title": "مجلس الجامعة يوافق على قبول التبرع لدعم مركز التميز في الرعاية الصحية لكلى الأطفال",
        "desc": "وافق مجلس جامعة الملك عبد العزيز على قبول التبرع العيني المقدم من أوقاف الشيخ عبد الرحمن عبد العزيز الراجحي، والمتمثل في توفير (9) أجهزة لغسيل الكلى مع محطة مياه متنقلة واحدة، لصالح مركز التميز في الرعاية الصحية لكلى الأطفال التابع للمستشفى الجامعي.",
    },
    {
        "img": static("images/news-media-lab.jpg"),
        "tag": "شراكة", "date": "",
        "title": "موافقة مجلس الجامعة على قبول التبرع العيني لصالح مشروع المعمل الإعلامي",
        "desc": "وافق مجلس الجامعة على قبول التبرع العيني المقدم من مؤسسة جرير للتسويق لدعم مشروع المعمل الإعلامي بكلية الاتصال والإعلام، وذلك في إطار تعزيز البيئة التعليمية والتدريبية، والارتقاء بالإمكانات التقنية التي تدعم العملية الأكاديمية.",
    },
    {
        "img": static("images/news-savvy-games.jpg"),
        "tag": "شراكة", "date": "",
        "title": "إبرام اتفاقية تعاون مع شركة الاستثمارات المحلية للألعاب الإلكترونية",
        "desc": "وافق مجلس الجامعة على إبرام اتفاقية تعاون بين الجامعة وشركة الاستثمارات المحلية للألعاب الإلكترونية (مجموعة سافي للألعاب)، وذلك في إطار تعزيز الشراكات الاستراتيجية مع القطاع الخاص، ودعم البرامج الأكاديمية المتخصصة التي تواكب متطلبات القطاعات الواعدة.",
    },
    {
        "img": static("images/news-children-voices.jpg"),
        "tag": "شراكة", "date": "",
        "title": 'مشروع "منهجية أصوات الأطفال" بمنحة من مؤسسة "أحياها" الإنسانية',
        "desc": 'أنجزت جامعة الملك عبد العزيز، ممثلة بكلية علوم الإنسان والتصاميم وقسم الطفولة المبكرة، المرحلة الأولى من مشروع "منهجية أصوات الأطفال"، الذي نفذ بدعم من مؤسسة الأمير عبد العزيز بن طلال وسُرى بنت سعود للتنمية الإنسانية (أحياها الإنسانية)، وبالشراكة مع جامعة كلية لندن الجامعية (UCL) ومؤسسة البداية المشرقة (Bright Start Foundation).',
    },
]

RESOURCE_CARDS = [
    {"img": "https://kau.edu.sa/web/static?path=uploads/2025/Jul/sections/REf9V0RFpkp7J06a.jfif", "title": "الخدمات"},
    {"img": "https://kau.edu.sa/web/static?path=uploads/2025/Jul/sections/gMmgcNnyngIHZo6J.jfif", "title": "السياسات والإجراءات"},
    {"img": "https://kau.edu.sa/web/static?path=uploads/2025/Jul/sections/40q7N5PKOjRcvRTh.jfif", "title": "النماذج والأدلة"},
    {"img": "https://kau.edu.sa/web/static?path=uploads/2025/Jul/sections/2qny7AcSHlBn8Rmx.jfif", "title": "دليل التواصل"},
]

CONTACT_ITEMS = [
    {"label": "رقم الهاتف", "value": "0126951328"},
    {"label": "البريد الإلكتروني", "value": "dra@kau.edu.sa"},
    {"label": "موقع الإدارة", "value": "www.dra.kau.edu.sa"},
]

PARTNERS = [
    {"src": static('images/partner-01-acino.jpeg'), "alt": 'Acino', "aspect": 'wide', "label": None},
    {"src": static('images/partner-02-partner-02.jpg'), "alt": 'Alrajhi Endowments', "aspect": 'wide', "label": None},
    {"src": static('images/partner-03-partner-03.png'), "alt": 'Al Dhuwayan Charity', "aspect": 'wide', "label": None},
    {"src": static('images/partner-04-partner-04.png'), "alt": 'Damy', "aspect": 'wide', "label": None},
    {"src": static('images/partner-05-partner-05.png'), "alt": 'The National Home Health Care Foundation', "aspect": 'wide', "label": None},
    {"src": static('images/partner-06-partner-06.png'), "alt": 'Al Dahyan Endowments', "aspect": 'tall', "label": None},
    {"src": static('images/partner-07-partner-07.png'), "alt": 'Saudi Orthodontic Society', "aspect": 'square', "label": None},
    {"src": static('images/partner-08-partner-08.jpeg'), "alt": 'Afaq Educational Association', "aspect": 'square', "label": None},
    {"src": static('images/partner-09-partner-09.jpg'), "alt": 'Safe Medical Care Org', "aspect": 'square', "label": None},
    {"src": static('images/partner-10-partner-10.png'), "alt": 'Dr. Abeer Mohammed Nuwaisir', "aspect": 'wide', "label": 'د. عبير محمد نويصر'},
    {"src": static('images/partner-11-partner-11.jpeg'), "alt": "Ta'alum", "aspect": 'square', "label": None},
    {"src": static('images/partner-12-partner-12.png'), "alt": 'Heart Patients Care Society', "aspect": 'square', "label": None},
    {"src": static('images/partner-13-partner-13.png'), "alt": 'Al Amjaad Trading & Manufacturing Company', "aspect": 'wide', "label": None},
    {"src": static('images/partner-14-partner-14.png'), "alt": 'Al Tayyar International Travel & Tourism Agency', "aspect": 'square', "label": None},
    {"src": static('images/partner-15-partner-15.png'), "alt": 'Baazeem Trading Co.', "aspect": 'wide', "label": None},
    {"src": static('images/partner-16-partner-16.jpeg'), "alt": 'Raz for Training Consultants', "aspect": 'square', "label": None},
    {"src": static('images/partner-17-sela.png'), "alt": 'Sela', "aspect": 'wide', "label": None},
    {"src": static('images/partner-18-keysight.png'), "alt": 'Keysight', "aspect": 'wide', "label": None},
    {"src": static('images/partner-19-savvy-games-group.png'), "alt": 'Savvy Games Group', "aspect": 'square', "label": None},
    {"src": static('images/partner-20-partner-20.png'), "alt": 'Ibrahim Al-Angary and His Descendants Foundation', "aspect": 'wide', "label": None},
    {"src": static('images/partner-21-partner-21.png'), "alt": 'Ahyaha', "aspect": 'wide', "label": None},
    {"src": static('images/partner-22-partner-22.png'), "alt": 'Hassan Taher Charity Foundation', "aspect": 'wide', "label": None},
    {"src": static('images/partner-23-partner-23.png'), "alt": 'Salem Bin Mahfouz Foundation', "aspect": 'wide', "label": None},
    {"src": static('images/partner-24-partner-24.png'), "alt": 'Berr Charity Fund - SABIC Employees', "aspect": 'wide', "label": None},
    {"src": static('images/partner-25-partner-25.png'), "alt": 'Bin Deajim', "aspect": 'wide', "label": None},
    {"src": static('images/partner-26-partner-26.png'), "alt": 'Al Subaie Charity Foundation', "aspect": 'wide', "label": None},
    {"src": static('images/partner-27-partner-27.png'), "alt": 'Ghaith', "aspect": 'wide', "label": None},
    {"src": static('images/partner-28-partner-28.png'), "alt": 'Nestlé', "aspect": 'wide', "label": None},
    {"src": static('images/partner-29-partner-29.png'), "alt": 'Allulwah Foundation', "aspect": 'wide', "label": None},
    {"src": static('images/partner-30-partner-30.png'), "alt": 'SABIC', "aspect": 'wide', "label": None},
    {"src": static('images/partner-31-partner-31.jpg'), "alt": 'Saudi Aramco', "aspect": 'wide', "label": None},
    {"src": static('images/partner-32-partner-32.png'), "alt": 'Dallah Albaraka', "aspect": 'wide', "label": None},
    {"src": static('images/partner-33-partner-33.png'), "alt": 'Baraa Society of Blood and Cancer in Makkah Children', "aspect": 'wide', "label": None},
    {"src": static('images/partner-34-partner-34.png'), "alt": 'Abdullah Taha Nejaish Endowment', "aspect": 'wide', "label": None},
    {"src": static('images/partner-35-partner-35.png'), "alt": 'Abudawood Group', "aspect": 'wide', "label": None},
    {"src": static('images/partner-36-partner-36.png'), "alt": 'Saudi National Bank', "aspect": 'wide', "label": None},
    {"src": static('images/partner-37-partner-37.jpeg'), "alt": 'Islamic Development Bank', "aspect": 'wide', "label": None},
    {"src": static('images/partner-38-partner-38.png'), "alt": 'Heart Friends Charitable Society', "aspect": 'wide', "label": None},
    {"src": static('images/partner-39-partner-39.jpeg'), "alt": 'Tamr', "aspect": 'wide', "label": None},
    {"src": static('images/partner-40-partner-40.jpg'), "alt": "Falathat Society for Children's Health Care", "aspect": 'wide', "label": None},
    {"src": static('images/partner-41-partner-41.jpg'), "alt": 'Hassan Abbas Sharbatly Community Service Foundation', "aspect": 'wide', "label": None},
    {"src": static('images/partner-42-partner-42.png'), "alt": 'Haif Company', "aspect": 'wide', "label": None},
    {"src": static('images/partner-43-partner-43.jpeg'), "alt": 'Sheikh Saleh Hamza Sairafi Endowment', "aspect": 'wide', "label": None},
    {"src": static('images/partner-44-partner-44.png'), "alt": 'KAU Scientific Endowment', "aspect": 'wide', "label": None},
    {"src": static('images/partner-45-partner-45.jpeg'), "alt": 'Hoshan Group', "aspect": 'wide', "label": None},
    {"src": static('images/partner-46-partner-46.png'), "alt": 'Jarir Bookstore', "aspect": 'wide', "label": None},
    {"src": static('images/partner-47-partner-47.png'), "alt": 'Savola Group', "aspect": 'wide', "label": None},
    {"src": static('images/partner-48-partner-48.png'), "alt": 'Jameel Community Saudi Arabia', "aspect": 'wide', "label": None},
    {"src": static('images/partner-49-partner-49.jpg'), "alt": 'Alagil Foundation', "aspect": 'wide', "label": None},
    {"src": None, "alt": 'وقف المدرسة الكشميرية', "aspect": 'wide', "label": 'وقف المدرسة الكشميرية'},
    {"src": None, "alt": 'الشيخ محمد حسين العامودي', "aspect": 'wide', "label": 'الشيخ محمد حسين العامودي'},
    {"src": None, "alt": 'الأمير نواف بن فيصل بن فهد آل سعود', "aspect": 'wide', "label": 'الأمير نواف بن فيصل بن فهد آل سعود'},
    {"src": None, "alt": 'الأمير بندر بن سلمان بن محمد آل سعود', "aspect": 'wide', "label": 'الأمير بندر بن سلمان بن محمد آل سعود'},
    {"src": None, "alt": 'سمو الأميرة الجوهرة الإبراهيم', "aspect": 'wide', "label": 'سمو الأميرة الجوهرة الإبراهيم'},
    {"src": None, "alt": 'الشيخ سالم صالح باعظيم', "aspect": 'wide', "label": 'الشيخ سالم صالح باعظيم'},
    {"src": None, "alt": 'الأمير خالد الفيصل آل سعود', "aspect": 'wide', "label": 'الأمير خالد الفيصل آل سعود'},
    {"src": None, "alt": 'أ.د. عصام أزهر', "aspect": 'wide', "label": 'أ.د. عصام أزهر'},
]


def home(request):
    context = {
        "program_steps": PROGRAM_STEPS,
        "goals": GOALS,
        "news_items_loop": (NEWS_ITEMS * 2),
        "resource_cards": RESOURCE_CARDS,
        "partners": PARTNERS,
        "partners_cols": 10,
        "partners_gap_total": 90,
        "contact_items": CONTACT_ITEMS,
    }
    return render(request, "pages/home.html", context)


# ─────────────────────────────────────────────────────────────
# صفحة "قبول المنح"
# ─────────────────────────────────────────────────────────────

GRANT_CONTACT_ITEMS = [
    {"label": "الهاتف", "value": "0126951328"},
    {"label": "البريد الإلكتروني", "value": "dra@kau.edu.sa"},
    {"label": "الموقع الإلكتروني", "value": "www.dra.kau.edu.sa"},
]


def grant_acceptance(request):
    context = {
        "program_steps": PROGRAM_STEPS,
        "contact_items": GRANT_CONTACT_ITEMS,
    }
    return render(request, "pages/grant_acceptance.html", context)


# ─────────────────────────────────────────────────────────────
# صفحة "اتصل بنا"
# ─────────────────────────────────────────────────────────────

SUGGESTIONS_CONTACT_ITEMS = [
    {"label": "رقم الهاتف", "value": "0126951328"},
    {"label": "البريد الإلكتروني", "value": "dra@kau.edu.sa"},
    {"label": "موقع الإدارة", "value": "www.dra.kau.edu.sa"},
]


def suggestions(request):
    context = {
        "contact_items": SUGGESTIONS_CONTACT_ITEMS,
    }
    return render(request, "pages/suggestions.html", context)


# ─────────────────────────────────────────────────────────────
# إرسال إيميل موحّد (نص + HTML) لكل نماذج الموقع
# ─────────────────────────────────────────────────────────────

def send_request_email(subject, rows):
    plain_body = "\n".join(f"{row['label']}: {row['value']}" for row in rows)
    html_body = render_to_string("emails/request_email.html", {"title": subject, "rows": rows})

    msg = EmailMultiAlternatives(
        subject=subject,
        body=plain_body,
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=[settings.GRANT_REQUESTS_RECIPIENT],
    )
    msg.attach_alternative(html_body, "text/html")
    msg.send(fail_silently=False)


GRANT_EXCEL_PATH = os.path.join(settings.MEDIA_ROOT, "exports", "grant_requests.xlsx")


def append_grant_to_excel(row_dict):
    os.makedirs(os.path.dirname(GRANT_EXCEL_PATH), exist_ok=True)
    if os.path.exists(GRANT_EXCEL_PATH):
        wb = load_workbook(GRANT_EXCEL_PATH)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "طلبات تقديم منح"
        ws.append(list(row_dict.keys()))
    ws.append(list(row_dict.values()))
    wb.save(GRANT_EXCEL_PATH)


def submit_grant(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "invalid_method"}, status=405)

    donor_name = request.POST.get("donor_name", "")
    phone = request.POST.get("phone", "")
    donation_type = request.POST.get("donation_type", "")
    email = request.POST.get("email", "")
    notes = request.POST.get("notes", "")

    rows = [
        {"label": "الاسم / الجهة المتبرعة", "value": donor_name},
        {"label": "رقم التواصل", "value": phone},
        {"label": "نوع التبرع", "value": donation_type},
        {"label": "البريد الإلكتروني", "value": email},
        {"label": "ملاحظات", "value": notes},
    ]

    try:
        send_request_email("طلب تقديم منحة جديد", rows)
        append_grant_to_excel({
            "التاريخ": timezone.localtime(timezone.now()).strftime("%Y-%m-%d %H:%M"),
            "الاسم / الجهة": donor_name,
            "رقم التواصل": phone,
            "نوع التبرع": donation_type,
            "البريد الإلكتروني": email,
            "ملاحظات": notes,
        })
        return JsonResponse({"success": True})
    except Exception:
        return JsonResponse({"success": False, "error": "send_failed"}, status=500)


def submit_suggestion(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "invalid_method"}, status=405)

    form_type = request.POST.get("form_type", "")

    if form_type == "external":
        subject = "اقتراح جديد - جهة من خارج الجامعة"
        rows = [
            {"label": "الاسم", "value": request.POST.get("name", "")},
            {"label": "اسم الجهة", "value": request.POST.get("org_name", "")},
            {"label": "رقم التواصل", "value": request.POST.get("phone", "")},
            {"label": "البريد الإلكتروني", "value": request.POST.get("email", "")},
            {"label": "الاقتراح", "value": request.POST.get("suggestion", "")},
        ]

    elif form_type == "internal":
        subject = "اقتراح جديد - جهة من داخل الجامعة"
        rows = [
            {"label": "اسم الجهة", "value": request.POST.get("org_name", "")},
            {"label": "اسم المشروع", "value": request.POST.get("project_name", "")},
            {"label": "فكرة المشروع", "value": request.POST.get("idea", "")},
            {"label": "الفئة المستهدفة", "value": request.POST.get("target_group", "")},
            {"label": "الطاقة الاستيعابية", "value": request.POST.get("capacity", "")},
            {"label": "مخرجات المشروع", "value": request.POST.get("outcomes", "")},
            {"label": "التكلفة التقديرية", "value": request.POST.get("estimated_cost", "")},
            {"label": "ضابط الاتصال", "value": request.POST.get("contact_officer", "")},
            {"label": "ملاحظات", "value": request.POST.get("notes", "")},
        ]

    elif form_type == "members":
        subject = "اقتراح جديد - منسوبين وطلبة"
        rows = [
            {"label": "الاسم", "value": request.POST.get("name", "")},
            {"label": "الرقم الجامعي/الوظيفي", "value": request.POST.get("id_number", "")},
            {"label": "الكلية أو القسم", "value": request.POST.get("department", "")},
            {"label": "رقم التواصل", "value": request.POST.get("phone", "")},
            {"label": "اسم المشروع", "value": request.POST.get("project_name", "")},
            {"label": "الفئة المستهدفة", "value": request.POST.get("target_group", "")},
            {"label": "فكرة المشروع", "value": request.POST.get("idea", "")},
            {"label": "مخرجات المشروع", "value": request.POST.get("outcomes", "")},
            {"label": "ملاحظات", "value": request.POST.get("notes", "")},
        ]

    else:
        return JsonResponse({"success": False, "error": "invalid_form_type"}, status=400)

    try:
        send_request_email(subject, rows)
        return JsonResponse({"success": True})
    except Exception:
        return JsonResponse({"success": False, "error": "send_failed"}, status=500)